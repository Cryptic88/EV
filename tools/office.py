"""Word and Excel file workflows; edits keep backups and preserve documents."""
import json
import os
import re


class OfficeTools:
    def __init__(self, files):
        self.files = files

    def word(self, action, path, text='', old_text='', new_text='', title=''):
        from docx import Document
        target = self.files.path(path, {'.docx'})
        if action == 'open':
            if not target.is_file():
                raise ValueError('Create or choose an existing Word document first.')
            os.startfile(str(target))
            return f'Opened {target} in its registered Word application.'
        if action == 'create':
            if target.exists():
                raise ValueError('Document exists. Use append or replace instead.')
            doc = Document()
            if title:
                doc.add_heading(title, 0)
            for paragraph in text.split('\n'):
                doc.add_paragraph(paragraph)
        else:
            doc = Document(target)
            paragraphs = list(doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            if action == 'read':
                return '\n'.join(p.text for p in paragraphs)[:16000]
            if action == 'append':
                for paragraph in text.split('\n'):
                    doc.add_paragraph(paragraph)
            elif action == 'replace':
                matches = [p for p in paragraphs if old_text and old_text in p.text]
                if sum(p.text.count(old_text) for p in matches) != 1:
                    raise ValueError('Old text must match exactly once in one paragraph.')
                paragraph = matches[0]
                start = paragraph.text.index(old_text)
                end = start + len(old_text)
                offset = 0
                inserted = False
                for run in paragraph.runs:
                    original = run.text
                    left, right = offset, offset + len(original)
                    if right > start and left < end:
                        prefix = original[:max(0, start - left)]
                        suffix = original[max(0, end - left):] if right > end else ''
                        run.text = prefix + (new_text if not inserted else '') + suffix
                        inserted = True
                    offset = right
                if not inserted:
                    raise ValueError('This paragraph uses unsupported text fields; edit it in Word.')
            else:
                raise ValueError('Word actions: create, open, read, append, replace.')
        return self.files.save(target, lambda tmp: doc.save(tmp))

    def excel(self, action, path, sheet='Sheet', cell='A1', values=None, formula='', rows=20, columns=10):
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils.cell import coordinate_to_tuple
        from openpyxl.workbook.properties import CalcProperties
        target = self.files.path(path, {'.xlsx'})
        if action == 'open':
            if not target.is_file():
                raise ValueError('Create or choose an existing Excel workbook first.')
            os.startfile(str(target))
            return f'Opened {target} in its registered spreadsheet application.'
        if action == 'create':
            if target.exists():
                raise ValueError('Workbook exists; use write or formula.')
            book = Workbook()
            book.active.title = sheet
        else:
            book = load_workbook(target)
        try:
            if action == 'sheets':
                return json.dumps(book.sheetnames)
            if sheet not in book.sheetnames:
                if action == 'add_sheet':
                    book.create_sheet(sheet)
                else:
                    raise ValueError('Choose a sheet from: ' + ', '.join(book.sheetnames))
            elif action == 'add_sheet':
                raise ValueError('That sheet already exists.')
            tab = book[sheet]
            start_row, start_col = coordinate_to_tuple(cell)
            if action == 'read':
                data = [[tab.cell(r, c).value for c in range(start_col, start_col + min(30, max(1, int(columns))))]
                        for r in range(start_row, start_row + min(100, max(1, int(rows))))]
                return json.dumps(data, default=str)[:16000]
            if action == 'formula':
                if not formula.startswith('='):
                    raise ValueError('A formula must start with =.')
                values = [[formula]]
            elif action not in {'create', 'write', 'add_sheet'}:
                raise ValueError('Excel actions: create, sheets, read, write, formula, add_sheet.')
            if values is not None:
                if not isinstance(values, list) or any(not isinstance(row, list) for row in values):
                    raise ValueError('Values must be a two-dimensional array.')
                if sum(len(row) for row in values) > 2000:
                    raise ValueError('Limit each write to 2,000 cells.')
                for r, row in enumerate(values, start_row):
                    for c, value in enumerate(row, start_col):
                        if isinstance(value, str) and value.startswith('=') and re.search(
                                r'\||\[|WEBSERVICE\s*\(|RTD\s*\(', value, re.I):
                            raise ValueError('External-data and executable formulas are excluded.')
                        tab.cell(r, c).value = value
            book.calculation = CalcProperties(fullCalcOnLoad=True)
            result = self.files.save(target, lambda tmp: book.save(tmp))
            return result + (' Formulas recalculate when opened in Excel.' if action == 'formula' else '')
        finally:
            book.close()
