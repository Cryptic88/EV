import asyncio
import json
from pathlib import Path
import tempfile
import time
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch
from tools.automation_common import FileAccess, ActionConfirmation
from tools.workspace import WorkspaceTools
from tools.office import OfficeTools
from tools.browser import BrowserTools
from tools.whatsapp import WhatsAppTools
from tools.outlook import OutlookTools


class FileWorkflowTests(unittest.TestCase):
    def setUp(self):
        folder = Path(__file__).resolve().parents[1] / 'artifacts/test-workspaces'
        folder.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=folder)
        self.root = Path(self.temp.name)
        self.files = FileAccess([self.root])
        self.workspace = WorkspaceTools(self.files)
        self.office = OfficeTools(self.files)

    def tearDown(self):
        self.temp.cleanup()

    def test_paths_cannot_escape_or_read_credentials(self):
        for path in ['../outside.txt', '.env', '.git/config', 'secret.pem']:
            with self.subTest(path=path), self.assertRaises(ValueError):
                self.files.path(path)

    def test_exact_code_edit_makes_recoverable_backup(self):
        target = self.root / 'hello.py'
        target.write_text('print("hello")\n', encoding='utf-8')
        result = self.workspace.workspace('replace', str(target), old_text='hello', new_text='world')
        self.assertEqual(target.read_text(), 'print("world")\n')
        backups = list((self.root / '.ev-backups').glob('*hello.py'))
        self.assertEqual(backups[0].read_text(), 'print("hello")\n')
        self.assertIn('Backup:', result)

    def test_ambiguous_code_edit_is_rejected(self):
        target = self.root / 'hello.py'
        target.write_text('hello hello', encoding='utf-8')
        with self.assertRaises(ValueError):
            self.workspace.workspace('replace', str(target), old_text='hello', new_text='world')
        self.assertEqual(target.read_text(), 'hello hello')

    def test_word_replacement_across_runs_preserves_surrounding_text(self):
        from docx import Document
        target = self.root / 'report.docx'
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Start hello ').bold = True
        paragraph.add_run('world end').italic = True
        document.save(target)
        self.office.word('replace', str(target), old_text='hello world', new_text='there')
        result = Document(target)
        self.assertEqual(result.paragraphs[0].text, 'Start there end')
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertTrue(result.paragraphs[0].runs[1].italic)

    def test_word_create_and_append_roundtrip(self):
        path = str(self.root / 'notes.docx')
        self.office.word('create', path, title='Notes', text='First thought')
        self.office.word('append', path, text='Second thought')
        text = self.office.word('read', path)
        self.assertIn('First thought', text)
        self.assertIn('Second thought', text)

    def test_excel_values_formulas_and_backup_roundtrip(self):
        path = str(self.root / 'budget.xlsx')
        self.office.excel('create', path, values=[['Item', 'Cost'], ['A', 10], ['B', 20]])
        self.office.excel('formula', path, cell='B4', formula='=SUM(B2:B3)')
        data = json.loads(self.office.excel('read', path, rows=4, columns=2))
        self.assertEqual(data[1][1], 10)
        self.assertEqual(data[3][1], '=SUM(B2:B3)')
        self.assertTrue(list((self.root / '.ev-backups').glob('*budget.xlsx')))

    def test_external_excel_formula_does_not_overwrite_file(self):
        path = str(self.root / 'budget.xlsx')
        self.office.excel('create', path, values=[[10]])
        original = Path(path).read_bytes()
        with self.assertRaises(ValueError):
            self.office.excel('formula', path, formula='=WEBSERVICE("https://example.com")')
        self.assertEqual(Path(path).read_bytes(), original)

    def test_real_python_test_job_reports_verified_result(self):
        (self.root / 'test_example.py').write_text(
            'import unittest\nclass TestExample(unittest.TestCase):\n def test_answer(self): self.assertEqual(2+2,4)\n', encoding='utf-8')
        self.workspace.workspace('test', str(self.root), framework='python')
        job_id = next(iter(self.workspace.jobs))
        deadline = time.monotonic() + 15
        while self.workspace.jobs[job_id]['status'] == 'running' and time.monotonic() < deadline:
            time.sleep(0.05)
        result = json.loads(self.workspace.workspace('job_status', job_id=job_id))
        self.assertEqual(result['status'], 'passed', result)
        self.assertEqual(result['exit_code'], 0)


class ConfirmationTests(unittest.TestCase):
    def test_confirmation_requires_an_unambiguous_affirmative(self):
        state, listener, transcriber, speaker = Mock(), Mock(), Mock(), Mock()
        state.is_stopped.return_value = False
        gate = ActionConfirmation(state, listener, transcriber, speaker)
        for answer, expected in [('yes', True), ('confirm', True), ('yesterday', False),
                                 ('yes but do not send', False), ('no', False), ('', False)]:
            transcriber.transcribe.return_value = answer
            self.assertEqual(gate.confirm('Send the reviewed draft?'), expected, answer)


class BrowserTests(unittest.TestCase):
    def test_url_rejects_local_files_scripts_and_credentials(self):
        for url in ['file:///C:/secret', 'javascript:alert(1)', 'https://user:pass@example.com']:
            with self.assertRaises(ValueError):
                BrowserTools.validate_url(url)
        BrowserTools.validate_url('http://127.0.0.1:5000')

    def setup_browser(self):
        gate = Mock()
        browser = BrowserTools(gate)
        page, node = Mock(), Mock()
        page.url = 'https://example.com/'
        page.is_closed.return_value = False
        node.count.return_value = 1
        node.is_visible.return_value = node.is_enabled.return_value = True
        browser.pages['page'] = page
        browser.elements['button'] = {'page_id': 'page', 'locator': node, 'url': page.url,
            'label': 'Send', 'submit': True, 'at': time.monotonic()}
        return browser, gate, page, node

    def test_cancelled_browser_submission_does_not_click(self):
        browser, gate, page, node = self.setup_browser()
        gate.confirm.return_value = False
        self.assertIn('Cancelled', browser.browser('submit', page_id='page', element_id='button'))
        node.click.assert_not_called()

    def test_navigation_during_confirmation_does_not_click(self):
        browser, gate, page, node = self.setup_browser()
        def confirm(_):
            page.url = 'https://example.org/'
            return True
        gate.confirm.side_effect = confirm
        with self.assertRaises(ValueError):
            browser.browser('submit', page_id='page', element_id='button')
        node.click.assert_not_called()

    def test_password_entry_and_stale_references_are_rejected(self):
        browser, gate, page, node = self.setup_browser()
        node.get_attribute.return_value = 'password'
        with self.assertRaises(ValueError):
            browser.browser('fill', page_id='page', element_id='button', text='secret')
        node.fill.assert_not_called()
        browser.elements['button']['at'] -= 121
        with self.assertRaises(ValueError):
            browser.browser('click', page_id='page', element_id='button')


class MessageTests(unittest.TestCase):
    def whatsapp(self):
        gate = Mock()
        tools = WhatsAppTools(Mock(), gate)
        page, composer, send = Mock(), Mock(), Mock()
        tools._page = Mock(return_value=page)
        tools.page_id = 'page'
        tools._verify_chat = Mock()
        tools._one = Mock(side_effect=[composer, send])
        composer.inner_text.return_value = 'Hello'
        tools.pending = {'contact': 'Test recipient', 'text': 'Hello', 'page_id': 'page'}
        return tools, gate, composer, send

    def test_whatsapp_cancel_does_not_send(self):
        tools, gate, composer, send = self.whatsapp()
        gate.confirm.return_value = False
        self.assertIn('Cancelled', tools.whatsapp('send'))
        send.click.assert_not_called()
        self.assertIsNotNone(tools.pending)

    def test_whatsapp_changed_draft_during_confirmation_does_not_send(self):
        tools, gate, composer, send = self.whatsapp()
        composer.inner_text.side_effect = ['Hello', 'Changed message']
        gate.confirm.return_value = True
        with self.assertRaises(ValueError):
            tools.whatsapp('send')
        send.click.assert_not_called()

    def mail(self):
        gate = Mock()
        tools = OutlookTools(gate)
        mail = Mock()
        mail.To, mail.CC, mail.BCC = 'test@example.com', '', ''
        mail.Subject, mail.Body = 'Test subject', 'Hello'
        mail.EntryID, mail.Parent.StoreID = 'entry', 'store'
        mail.Recipients.Count, mail.Attachments.Count = 0, 0
        app = Mock()
        app.Session.GetItemFromID.return_value = mail
        tools._app = Mock(return_value=app)
        tools.references['draft'] = ('draft', 'entry', 'store')
        tools.drafts['draft'] = tools._mail_fingerprint(mail)
        return tools, gate, mail

    def test_outlook_cancel_does_not_send(self):
        tools, gate, mail = self.mail()
        gate.confirm.return_value = False
        self.assertIn('Cancelled', tools.outlook('send', item_id='draft'))
        mail.Send.assert_not_called()

    def test_outlook_approved_unchanged_draft_calls_send(self):
        tools, gate, mail = self.mail()
        gate.confirm.return_value = True
        result = tools.outlook('send', item_id='draft')
        mail.Send.assert_called_once()
        self.assertIn('Delivery is not confirmed', result)

    def test_outlook_attachments_require_manual_review(self):
        tools, gate, mail = self.mail()
        mail.Attachments.Count = 1
        with self.assertRaises(ValueError):
            tools.outlook('send', item_id='draft')
        mail.Send.assert_not_called()

    def test_whatsapp_approved_unchanged_draft_clicks_send(self):
        tools, gate, composer, send = self.whatsapp()
        gate.confirm.return_value = True
        self.assertIn('send requested', tools.whatsapp('send'))
        send.click.assert_called_once()
        self.assertIsNone(tools.pending)

    def test_outlook_modified_draft_does_not_send(self):
        tools, gate, mail = self.mail()
        mail.Body = 'Changed'
        with self.assertRaises(ValueError):
            tools.outlook('send', item_id='draft')
        mail.Send.assert_not_called()
        gate.confirm.assert_not_called()

    def test_outlook_change_during_confirmation_does_not_send(self):
        tools, gate, mail = self.mail()
        def confirm(_):
            mail.To = 'other@example.com'
            return True
        gate.confirm.side_effect = confirm
        with self.assertRaises(ValueError):
            tools.outlook('send', item_id='draft')
        mail.Send.assert_not_called()

    def test_calendar_cancel_does_not_save_or_delete(self):
        tools, gate, mail = self.mail()
        tools.references['event'] = ('event', 'entry', 'store')
        mail.IsRecurring, mail.MeetingStatus = False, 0
        gate.confirm.return_value = False
        tools.calendar('delete', item_id='event')
        mail.Delete.assert_not_called()
        tools.calendar('update', item_id='event', subject='Changed')
        mail.Save.assert_not_called()

    def test_calendar_interprets_naive_dates_in_user_timezone(self):
        date = OutlookTools._date('2026-10-16T09:00:00')
        self.assertEqual(date.utcoffset().total_seconds(), 7200)


if __name__ == '__main__':
    unittest.main()
